"""Excel (.xlsx) raporları (openpyxl)."""

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASLIK_DOLGU = PatternFill('solid', fgColor='E9ECEF')
INCE = Side(style='thin')
KENARLIK = Border(top=INCE, left=INCE, bottom=INCE, right=INCE)


def _bicimle(ws, baslik_satiri=1, genislikler=None):
    for row in ws.iter_rows():
        for cell in row:
            cell.border = KENARLIK
    for r in range(1, baslik_satiri + 1):
        for cell in ws[r]:
            cell.font = Font(bold=True)
            cell.fill = BASLIK_DOLGU
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, w in enumerate(genislikler or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=baslik_satiri + 1, column=1)


def _kaydet(wb):
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sonuc_raporu(veri, ozet):
    sorular = veri['sorular']
    wb = Workbook()

    # 1) Katılımcılar
    ws = wb.active
    ws.title = 'Katılımcılar'
    ws.append(['Sıra No', 'Ad Soyad', 'Sınıf', 'Veli Kodu', 'Tarih', 'Yanıtlar'])
    for i, s in enumerate(veri['satirlar'], start=1):
        y = s['yanit']
        ws.append([i, y.ad_soyad, y.sinif.ad, y.kod.kod if y.kod else '',
                   timezone.localtime(y.tarih).replace(tzinfo=None), ', '.join(s['secimler'])])
    for cell in ws['E'][1:]:
        cell.number_format = 'dd.mm.yyyy hh:mm'
    _bicimle(ws, genislikler=[9, 28, 12, 12, 18, 40])

    # 2) Soru bazlı dağılım
    ws = wb.create_sheet('Soru Bazlı Dağılım')
    ws.append(['Soru', 'A Seçeneği', 'A Sayısı', 'A %', 'B Seçeneği', 'B Sayısı', 'B %'])
    for st in veri['soru_istatistik']:
        s = st['soru']
        ws.append([f'{s.sira}. Soru', s.secenek_a, st['a'], st['yuzde_a'] / 100, s.secenek_b, st['b'], st['yuzde_b'] / 100])
    for col in ('D', 'G'):
        for cell in ws[col][1:]:
            cell.number_format = '0%'
    for col in ('B', 'E'):
        for cell in ws[col][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    _bicimle(ws, genislikler=[8, 50, 10, 8, 50, 10, 8])

    # 3) Bireysel yanıtlar (M1..Mn, A/B alt sütunları)
    ws = wb.create_sheet('Bireysel Yanıtlar')
    ust = ['No', 'Adı Soyadı', 'Sınıf']
    alt = ['', '', '']
    for i in range(len(sorular)):
        ust += [f'M{i + 1}', '']
        alt += ['A', 'B']
    ws.append(ust)
    ws.append(alt)
    for c in (1, 2, 3):
        ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
    for i in range(len(sorular)):
        ws.merge_cells(start_row=1, start_column=4 + i * 2, end_row=1, end_column=5 + i * 2)
    for i, s in enumerate(veri['satirlar'], start=1):
        satir = [i, s['yanit'].ad_soyad, s['yanit'].sinif.ad]
        for secim in s['secimler']:
            satir += [1 if secim == 'A' else None, 1 if secim == 'B' else None]
        ws.append(satir)
    toplam = ['', 'TOPLAM', '']
    for st in veri['soru_istatistik']:
        toplam += [st['a'], st['b']]
    ws.append(toplam)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    _bicimle(ws, baslik_satiri=2, genislikler=[6, 26, 10] + [4.5] * (len(sorular) * 2))
    for row in ws.iter_rows(min_row=3, min_col=4):
        for cell in row:
            cell.alignment = Alignment(horizontal='center')

    # 4) Şık bazlı sonuçlar
    ws = wb.create_sheet('Şık Bazlı Sonuçlar')
    ws.append(['Soru', 'Şık', 'Seçen Sayısı', 'Yüzde'])
    for st in veri['soru_istatistik']:
        s = st['soru']
        ws.append([f'{s.sira}A', s.kisa_a, st['a'], st['yuzde_a'] / 100])
        ws.append([f'{s.sira}B', s.kisa_b, st['b'], st['yuzde_b'] / 100])
    for cell in ws['D'][1:]:
        cell.number_format = '0%'
    _bicimle(ws, genislikler=[8, 70, 14, 10])

    # 5) Sınıf özeti
    ws = wb.create_sheet('Sınıf Özeti')
    ws.append(['Sınıf', 'Katılan Veli', 'Dağıtılan Kod', 'Katılım Oranı'])
    for o in ozet:
        ws.append([o['sinif'].ad, o['katilan'], o['kod_sayisi'], o['oran'] / 100 if o['kod_sayisi'] else None])
    for cell in ws['D'][1:]:
        cell.number_format = '0%'
    _bicimle(ws, genislikler=[14, 14, 14, 14])

    return _kaydet(wb)


def kod_raporu(kodlar):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Veli Kodları'
    ws.append(['Sınıf', 'Veli Kodu', 'Durum', 'Kullanan Veli'])
    for k in kodlar:
        yanit = getattr(k, 'yanit', None)
        ws.append([k.sinif.ad, k.kod, 'Kullanıldı' if k.kullanildi else 'Kullanılmadı', yanit.ad_soyad if yanit else ''])
    _bicimle(ws, genislikler=[14, 14, 14, 28])
    return _kaydet(wb)
