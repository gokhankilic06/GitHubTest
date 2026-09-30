from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .models import Sinif, Soru, VeliYaniti


def cevaplar(secim):
    return {f'soru_{s.id}': secim for s in Soru.objects.filter(aktif=True)}


class AnketTesti(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser('yonetici', 'y@example.com', 'sifre-12345')

    def gonder(self, ad='Ali Veli', sinif='1-A', secim='A', **ekstra):
        return self.client.post(reverse('anket:anket'), {'ad_soyad': ad, 'sinif': sinif, **cevaplar(secim), **ekstra})

    def test_baslangic_verisi(self):
        self.assertEqual(Soru.objects.count(), 13)
        self.assertEqual(Sinif.objects.count(), 22)

    def test_anket_formu_sinif_listesi(self):
        r = self.client.get(reverse('anket:anket'))
        self.assertContains(r, 'name="sinif"')
        self.assertContains(r, '<option value="1-A"')
        self.assertContains(r, 'class="question"', count=13)

    def test_anket_gonderimi(self):
        r = self.gonder()
        self.assertRedirects(r, reverse('anket:tesekkurler'))
        yanit = VeliYaniti.objects.get()
        self.assertEqual(yanit.sinif.ad, '1-A')
        self.assertEqual(yanit.ad_soyad, 'Ali Veli')
        self.assertEqual(yanit.cevaplar.count(), 13)

    def test_eksik_bilgi_reddedilir_ve_secimler_korunur(self):
        veri = {'ad_soyad': 'Ali', 'sinif': '2-B', **cevaplar('B')}
        ilk = Soru.objects.get(sira=1)
        del veri[f'soru_{ilk.id}']
        r = self.client.post(reverse('anket:anket'), veri)
        self.assertEqual(r.status_code, 400)
        self.assertContains(r, 'tüm soruları yanıtlayınız', status_code=400)
        self.assertContains(r, 'class="question missing"', count=1, status_code=400)
        self.assertContains(r, 'value="B" checked', count=12, status_code=400)
        self.assertContains(r, '<option value="2-B" selected>', status_code=400)
        self.assertEqual(VeliYaniti.objects.count(), 0)

        self.assertEqual(self.gonder(ad='  ').status_code, 400)
        self.assertEqual(self.gonder(sinif='').status_code, 400)
        self.assertEqual(self.gonder(sinif='Olmayan').status_code, 400)
        self.assertEqual(VeliYaniti.objects.count(), 0)

    def test_panel_giris_gerektirir(self):
        for ad in ['anket:panel', 'anket:sonuc_excel']:
            r = self.client.get(reverse(ad))
            self.assertEqual(r.status_code, 302, ad)
            self.assertIn('/admin/login/', r['Location'])

    def test_panel_sinif_filtresi_ve_istatistik(self):
        self.gonder('Ali', '1-A', 'A')
        self.gonder('Ayşe', '1-A', 'B')
        self.gonder('Mehmet', '2-B', 'B')
        self.client.force_login(self.admin)

        r = self.client.get(reverse('anket:panel'))
        self.assertEqual(r.context['toplam'], 3)
        ilk = r.context['soru_istatistik'][0]
        self.assertEqual((ilk['a'], ilk['b'], ilk['yuzde_a'], ilk['yuzde_b']), (1, 2, 33, 67))
        ozet = {o['sinif'].ad: o['katilan'] for o in r.context['ozet']}
        self.assertEqual(ozet, {'1-A': 2, '2-B': 1})
        self.assertContains(r, '3 kişi ankete katılmıştır')

        r = self.client.get(reverse('anket:panel'), {'sinif': '1-A'})
        self.assertEqual(r.context['toplam'], 2)
        ilk = r.context['soru_istatistik'][0]
        self.assertEqual((ilk['a'], ilk['b'], ilk['yuzde_a']), (1, 1, 50))
        self.assertContains(r, '1-A sınıfından 2 kişi')

    def test_excel_raporu(self):
        self.gonder('Ali', '1-A', 'A')
        self.gonder('Mehmet', '2-B', 'B')
        self.client.force_login(self.admin)

        r = self.client.get(reverse('anket:sonuc_excel'), {'sinif': '1-A'})
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['Katılımcılar', 'Soru Bazlı Dağılım', 'Bireysel Yanıtlar', 'Şık Bazlı Sonuçlar', 'Sınıf Özeti'])
        katilim = wb['Katılımcılar']
        self.assertEqual(katilim.max_row, 2)
        self.assertEqual(katilim['C2'].value, '1-A')
        self.assertEqual(wb['Şık Bazlı Sonuçlar'].max_row, 1 + 26)
        self.assertEqual(wb['Sınıf Özeti'].max_row, 3)

    def test_yanit_silme(self):
        self.gonder()
        self.client.force_login(self.admin)
        yanit = VeliYaniti.objects.get()
        r = self.client.post(reverse('anket:yanit_sil', args=[yanit.pk]))
        self.assertEqual(r.status_code, 302)
        self.assertFalse(VeliYaniti.objects.exists())

    def test_sifirla(self):
        self.gonder()
        self.gonder('Ayşe')
        self.client.force_login(self.admin)
        self.client.post(reverse('anket:sifirla'))
        self.assertFalse(VeliYaniti.objects.exists())

    def test_pasif_soru_ankette_gorunmez(self):
        Soru.objects.filter(sira=13).update(aktif=False)
        r = self.client.get(reverse('anket:anket'))
        self.assertContains(r, 'class="question"', count=12)
        self.assertRedirects(self.gonder(), reverse('anket:tesekkurler'))
